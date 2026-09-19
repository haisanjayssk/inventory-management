import io
import csv
import json
from datetime import datetime, timezone
from pymongo import InsertOne, UpdateOne

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    openpyxl = None
    OPENPYXL_AVAILABLE = False

from app.repositories.item_type_repository import ItemTypeRepository
from app.repositories.item_repository import ItemRepository
from app.repositories.vendor_repository import VendorRepository
from app.models.counter import SequenceCounter

class ItemService:
    def __init__(self):
        self.item_type_repo = ItemTypeRepository()
        self.item_repo = ItemRepository()
        self.vendor_repo = VendorRepository()

    # --- Item Types ---
    def get_all_item_types(self, org_id: str = "ORG-001"):
        return self.item_type_repo.find_all_by_org(org_id)

    def get_item_type_by_id(self, item_type_id: str, org_id: str = "ORG-001"):
        it = self.item_type_repo.find_one({"_id": item_type_id, "organization_id": org_id})
        if not it:
            it = self.item_type_repo.find_by_code(org_id, item_type_id)
        if not it:
            raise ValueError(f"Item Type '{item_type_id}' not found")
        return it

    def create_item_type(self, data: dict, org_id: str = "ORG-001"):
        code = data["code"].strip().upper()
        if self.item_type_repo.find_by_code(org_id, code):
            raise ValueError(f"Item Type code '{code}' already exists")

        it_id = SequenceCounter.get_next_id("item_type")
        now = datetime.now(timezone.utc).isoformat()
        doc = {
            "_id": it_id,
            "organization_id": org_id,
            "code": code,
            "name": data["name"].strip(),
            "part_type_name": data["name"].strip(),
            "description": data.get("description", ""),
            "tracking_mode": data.get("tracking_mode", "QUANTITY"),
            "fields": data.get("fields", []),
            "site_id": data.get("site_id"),
            "active": data.get("active", True),
            "created_at": now,
            "updated_at": now
        }
        self.item_type_repo.insert_one(doc)
        return doc

    def update_item_type(self, item_type_id: str, data: dict, org_id: str = "ORG-001"):
        it = self.get_item_type_by_id(item_type_id, org_id)
        update_data = {k: v for k, v in data.items() if k not in ["_id", "organization_id", "created_at"]}
        if "name" in update_data:
            update_data["part_type_name"] = update_data["name"]
        update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
        self.item_type_repo.update_one({"_id": it["_id"]}, {"$set": update_data})
        return self.get_item_type_by_id(it["_id"], org_id)

    # --- Items ---
    def get_all_items(self, org_id: str = "ORG-001", item_type_id: str = None, search: str = None, active_only: bool = False):
        query = {"organization_id": org_id}
        if item_type_id:
            query["$or"] = [{"item_type_id": item_type_id}, {"part_type_id": item_type_id}]
        if active_only:
            query["active"] = True
        if search:
            s_clean = search.strip()
            query["$or"] = [
                {"code": {"$regex": s_clean, "$options": "i"}},
                {"part_code": {"$regex": s_clean, "$options": "i"}},
                {"name": {"$regex": s_clean, "$options": "i"}},
                {"part_name": {"$regex": s_clean, "$options": "i"}},
                {"description": {"$regex": s_clean, "$options": "i"}}
            ]
        items = self.item_repo.find_all(query, sort_by=[("code", 1)])
        types_cache = {t["_id"]: t for t in self.item_type_repo.find_all_by_org(org_id)}
        for item in items:
            it = types_cache.get(item.get("item_type_id"))
            if it:
                item["item_type_name"] = it.get("name")
                item["part_type_name"] = it.get("name")
                item["tracking_mode"] = it.get("tracking_mode", "QUANTITY")
                item["tracking_type"] = it.get("tracking_mode", "QUANTITY")
            # Populate legacy attributes
            item["part_code"] = item.get("code")
            item["part_name"] = item.get("name")
            item["part_id"] = item["_id"]
        return items

    def get_item_by_id(self, item_id: str, org_id: str = "ORG-001"):
        item = self.item_repo.find_one({"_id": item_id, "organization_id": org_id})
        if not item:
            item = self.item_repo.find_by_code(org_id, item_id)
        if not item:
            raise ValueError(f"Item '{item_id}' not found")

        it = self.item_type_repo.find_one({"_id": item.get("item_type_id"), "organization_id": org_id})
        if it:
            item["item_type_name"] = it.get("name")
            item["part_type_name"] = it.get("name")
            item["tracking_mode"] = it.get("tracking_mode", "QUANTITY")
            item["tracking_type"] = it.get("tracking_mode", "QUANTITY")
        item["part_code"] = item.get("code")
        item["part_name"] = item.get("name")
        item["part_id"] = item["_id"]
        return item

    def _validate_dynamic_attributes(self, item_type: dict, attributes: dict):
        errors = []
        fields = item_type.get("fields", [])
        for f in fields:
            f_key = f.get("field_key") or f.get("key")
            f_name = f.get("field_name") or f.get("name") or f_key
            if not f_key:
                continue
            req = f.get("required", False)
            d_type = (f.get("data_type") or f.get("type") or "STRING").upper()
            val = attributes.get(f_key)

            if req and (val is None or str(val).strip() == ""):
                errors.append(f"Field '{f_name}' ({f_key}) is required for type '{item_type.get('name')}'.")
                continue

            if val is not None and str(val).strip() != "":
                if d_type in ["NUMBER", "INTEGER"]:
                    try:
                        int(val)
                    except ValueError:
                        errors.append(f"Field '{f_name}' must be an integer.")
                elif d_type == "DECIMAL":
                    try:
                        float(val)
                    except ValueError:
                        errors.append(f"Field '{f_name}' must be a numeric/decimal value.")
                elif d_type == "BOOLEAN":
                    if not isinstance(val, bool) and str(val).lower() not in ["true", "false", "1", "0"]:
                        errors.append(f"Field '{f_name}' must be a boolean.")
                elif d_type == "SELECT":
                    opts = f.get("options", [])
                    if opts and str(val) not in opts:
                        errors.append(f"Field '{f_name}' value '{val}' is not in valid options: {opts}.")

        if errors:
            raise ValueError("; ".join(errors))

    def create_item(self, data: dict, org_id: str = "ORG-001"):
        code = (data.get("code") or data.get("part_code") or "").strip().upper()
        name = (data.get("name") or data.get("part_name") or "").strip()
        type_id = (data.get("item_type_id") or data.get("part_type_id") or "").strip()

        if self.item_repo.find_by_code(org_id, code):
            raise ValueError(f"Item code '{code}' already exists")

        it = self.get_item_type_by_id(type_id, org_id)
        attributes = data.get("attributes", {})
        self._validate_dynamic_attributes(it, attributes)

        item_id = SequenceCounter.get_next_id("item")
        now = datetime.now(timezone.utc).isoformat()
        doc = {
            "_id": item_id,
            "organization_id": org_id,
            "item_type_id": it["_id"],
            "code": code,
            "part_code": code,
            "name": name,
            "part_name": name,
            "description": data.get("description", ""),
            "attributes": attributes,
            "package": data.get("package", ""),
            "mpn": data.get("mpn", ""),
            "mfr": data.get("mfr", ""),
            "vendor_id": data.get("vendor_id"),
            "site_id": data.get("site_id") or it.get("site_id") or "SITE-001",
            "active": data.get("active", True),
            "created_at": now,
            "updated_at": now
        }
        self.item_repo.insert_one(doc)
        return doc

    def update_item(self, item_id: str, data: dict, org_id: str = "ORG-001"):
        item = self.get_item_by_id(item_id, org_id)
        it = self.get_item_type_by_id(item["item_type_id"], org_id)
        attributes = data.get("attributes", item.get("attributes", {}))
        self._validate_dynamic_attributes(it, attributes)

        update_data = {k: v for k, v in data.items() if k not in ["_id", "organization_id", "code", "created_at"]}
        if "name" in update_data:
            update_data["part_name"] = update_data["name"]
        update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
        self.item_repo.update_one({"_id": item["_id"]}, {"$set": update_data})
        return self.get_item_by_id(item["_id"], org_id)

    def delete_item(self, item_id: str, org_id: str = "ORG-001"):
        item = self.get_item_by_id(item_id, org_id)
        self.item_repo.update_one({"_id": item["_id"]}, {"$set": {"active": False}})
        return {"item_id": item["_id"], "message": "Item deactivated successfully"}

    # --- Bulk Import / Export ---
    def generate_import_template(self, org_id: str = "ORG-001", file_format: str = "csv"):
        item_types = self.item_type_repo.find_all_by_org(org_id)
        dynamic_cols = []
        seen_keys = set()
        for it in item_types:
            for f in it.get("fields", []):
                key = f.get("field_key") or f.get("key")
                if key and key not in seen_keys:
                    seen_keys.add(key)
                    dynamic_cols.append(f"attr_{key}")

        standard_headers = ["part_code", "part_name", "part_type", "tracking_type", "description", "site_id", "vendor"]
        all_headers = standard_headers + dynamic_cols

        sample_rows = [
            {
                "part_code": "RES-10K-0603",
                "part_name": "10K Ohm SMD Resistor 0603",
                "part_type": "RESISTOR",
                "tracking_type": "QUANTITY",
                "description": "Thick film chip resistor",
                "site_id": "SITE-001",
                "vendor": "Yageo",
                "attr_resistance": "10000",
                "attr_tolerance": "1",
                "attr_power_rating": "0.1"
            },
            {
                "part_code": "CELL-21700-50E",
                "part_name": "21700 5000mAh Battery Cell",
                "part_type": "CELL",
                "tracking_type": "SERIAL",
                "description": "High energy cylindrical battery cell",
                "site_id": "SITE-001",
                "vendor": "Samsung SDI",
                "attr_nominal_voltage": "3.6",
                "attr_capacity": "5000",
                "attr_chemistry": "NMC"
            }
        ]

        if file_format.lower() in ["xlsx", "excel"]:
            if not OPENPYXL_AVAILABLE:
                raise ValueError("Excel template generation requires 'openpyxl'.")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Parts Catalog Template"
            header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            align_center = Alignment(horizontal="center", vertical="center")

            ws.append(all_headers)
            for col_idx in range(1, len(all_headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            for item in sample_rows:
                ws.append([item.get(h, "") for h in all_headers])

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "items_import_template.xlsx"
        else:
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=all_headers)
            writer.writeheader()
            for r in sample_rows:
                writer.writerow({h: r.get(h, "") for h in all_headers})
            return output.getvalue().encode("utf-8-sig"), "text/csv", "items_import_template.csv"

    def bulk_import_items(self, file_bytes: bytes, filename: str, org_id: str = "ORG-001", duplicate_strategy: str = "skip"):
        lower_name = (filename or "").lower()
        rows = []

        if lower_name.endswith(".json"):
            try:
                raw_json = json.loads(file_bytes.decode("utf-8"))
                rows = raw_json if isinstance(raw_json, list) else raw_json.get("items", [])
            except Exception as e:
                raise ValueError(f"Invalid JSON file format: {str(e)}")
        elif lower_name.endswith(".xlsx") or lower_name.endswith(".xls"):
            if not OPENPYXL_AVAILABLE:
                raise ValueError("Excel file reading requires 'openpyxl'.")
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            iter_rows = list(ws.iter_rows(values_only=True))
            if not iter_rows:
                raise ValueError("The Excel worksheet is empty.")
            headers = [str(h).strip() if h is not None else "" for h in iter_rows[0]]
            for r in iter_rows[1:]:
                if any(cell is not None and str(cell).strip() != "" for cell in r):
                    row_dict = {}
                    for idx, h in enumerate(headers):
                        if h:
                            row_dict[h] = r[idx] if idx < len(r) and r[idx] is not None else ""
                    rows.append(row_dict)
        else:
            text = None
            for encoding in ["utf-8-sig", "utf-8", "latin-1"]:
                try:
                    text = file_bytes.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ValueError("Unable to decode CSV file.")
            reader = csv.DictReader(io.StringIO(text))
            for r in reader:
                if any(val is not None and str(val).strip() != "" for val in r.values()):
                    cleaned = {k.strip(): str(v).strip() if v is not None else "" for k, v in r.items() if k}
                    rows.append(cleaned)

        if not rows:
            raise ValueError("No data rows found to import.")

        db_item_types = self.item_type_repo.find_all_by_org(org_id)
        types_by_id = {t["_id"].upper(): t for t in db_item_types}
        types_by_code = {t["code"].upper(): t for t in db_item_types}
        types_by_name = {t["name"].upper(): t for t in db_item_types}

        seen_codes = set()
        errors = []
        to_insert_docs = []
        to_update_ops = []
        auto_created_vendors = []
        skipped_count = 0
        now_str = datetime.now(timezone.utc).isoformat()

        for row_idx, row in enumerate(rows, start=2):
            row_norm = {k.strip().lower().replace(" ", "_"): v for k, v in row.items()}
            code = str(row_norm.get("code") or row_norm.get("part_code") or row_norm.get("item_code") or "").strip()
            if not code:
                errors.append({"row": row_idx, "code": "N/A", "error": "Item code is required."})
                continue

            code_upper = code.upper()
            if code_upper in seen_codes:
                errors.append({"row": row_idx, "code": code, "error": f"Duplicate code '{code}' in import file."})
                continue
            seen_codes.add(code_upper)

            name = str(row_norm.get("name") or row_norm.get("part_name") or row_norm.get("description") or "").strip()
            if not name:
                errors.append({"row": row_idx, "code": code, "error": "Item name is required."})
                continue

            type_val = str(row_norm.get("item_type") or row_norm.get("type") or row_norm.get("part_type") or "").strip().upper()
            target_type = types_by_id.get(type_val) or types_by_code.get(type_val) or types_by_name.get(type_val)
            if not target_type and len(db_item_types) == 1:
                target_type = db_item_types[0]

            if not target_type:
                errors.append({"row": row_idx, "code": code, "error": f"Item type '{type_val}' not found."})
                continue

            # Auto-create vendor if specified
            vendor_name_val = str(row_norm.get("vendor") or row_norm.get("vendor_name") or "").strip()
            vendor_id_val = None
            if vendor_name_val:
                existing_v = self.vendor_repo.find_by_name(org_id, vendor_name_val)
                if not existing_v:
                    new_vid = SequenceCounter.get_next_id("vendor")
                    v_doc = {
                        "_id": new_vid,
                        "organization_id": org_id,
                        "code": f"VEN-{new_vid[-3:] if len(new_vid)>=7 else '001'}",
                        "name": vendor_name_val,
                        "vendor_name": vendor_name_val,
                        "types": ["DISTRIBUTOR"],
                        "contact": {},
                        "address": {},
                        "site_id": str(row_norm.get("site_id") or "SITE-001"),
                        "status": "ACTIVE",
                        "created_at": now_str,
                        "updated_at": now_str
                    }
                    self.vendor_repo.insert_one(v_doc)
                    auto_created_vendors.append(vendor_name_val)
                    vendor_id_val = new_vid
                else:
                    vendor_id_val = existing_v["_id"]

            attributes = {}
            for field in target_type.get("fields", []):
                f_key = field.get("field_key") or field.get("key") or (field.get("field_name") or field.get("name", "")).strip().lower().replace(" ", "_")
                if not f_key:
                    continue
                val = row_norm.get(f"attr_{f_key}") or row_norm.get(f_key)
                if val is not None and str(val).strip() != "":
                    attributes[f_key] = val

            for k, v in row_norm.items():
                if k.startswith("attr_") and v is not None and str(v).strip() != "":
                    attributes[k[5:]] = str(v).strip()

            existing_item = self.item_repo.find_by_code(org_id, code_upper)
            if existing_item:
                if duplicate_strategy == "skip":
                    skipped_count += 1
                    continue
                elif duplicate_strategy == "error":
                    errors.append({"row": row_idx, "code": code, "error": f"Item code '{code}' already exists."})
                    continue
                elif duplicate_strategy == "update":
                    update_dict = {
                        "name": name,
                        "part_name": name,
                        "description": str(row_norm.get("description") or "").strip(),
                        "attributes": attributes,
                        "site_id": str(row_norm.get("site_id") or "").strip() or None,
                        "active": True,
                        "updated_at": now_str
                    }
                    if vendor_id_val:
                        update_dict["vendor_id"] = vendor_id_val
                        update_dict["vendor_name"] = vendor_name_val
                    to_update_ops.append(
                        UpdateOne(
                            {"_id": existing_item["_id"]},
                            {"$set": update_dict}
                        )
                    )
            else:
                new_item_doc = {
                    "organization_id": org_id,
                    "item_type_id": target_type["_id"],
                    "code": code_upper,
                    "part_code": code_upper,
                    "name": name,
                    "part_name": name,
                    "description": str(row_norm.get("description") or "").strip(),
                    "attributes": attributes,
                    "site_id": str(row_norm.get("site_id") or "").strip() or None,
                    "active": True
                }
                if vendor_id_val:
                    new_item_doc["vendor_id"] = vendor_id_val
                    new_item_doc["vendor_name"] = vendor_name_val
                to_insert_docs.append(new_item_doc)

        if to_insert_docs:
            batch_ids = SequenceCounter.get_next_batch_ids("item", len(to_insert_docs))
            for doc, i_id in zip(to_insert_docs, batch_ids):
                doc["_id"] = i_id
                doc["created_at"] = now_str
                doc["updated_at"] = now_str

        operations = []
        if to_insert_docs:
            operations.extend([InsertOne(d) for d in to_insert_docs])
        if to_update_ops:
            operations.extend(to_update_ops)

        if operations:
            self.item_repo.bulk_write(operations, ordered=False)

        return {
            "total_rows": len(rows),
            "imported": len(to_insert_docs),
            "updated": len(to_update_ops),
            "skipped": skipped_count,
            "errors": errors,
            "auto_created_vendors": list(dict.fromkeys(auto_created_vendors))
        }

    # Backward compatibility alias
    bulk_import_parts = bulk_import_items

    def export_items(self, org_id: str = "ORG-001", file_format: str = "csv"):
        items = self.get_all_items(org_id)
        item_types = self.item_type_repo.find_all_by_org(org_id)
        types_cache = {t["_id"]: t for t in item_types}

        dynamic_cols = []
        seen_keys = set()
        for it in item_types:
            for f in it.get("fields", []):
                key = f.get("field_key") or f.get("key")
                if key and key not in seen_keys:
                    seen_keys.add(key)
                    dynamic_cols.append(f"attr_{key}")

        standard_headers = ["code", "name", "item_type", "tracking_mode", "description", "site_id", "vendor_id"]
        all_headers = standard_headers + dynamic_cols

        rows = []
        for item in items:
            it = types_cache.get(item.get("item_type_id"))
            row = {
                "code": item.get("code", ""),
                "name": item.get("name", ""),
                "item_type": it.get("name", "") if it else "",
                "tracking_mode": it.get("tracking_mode", "QUANTITY") if it else "QUANTITY",
                "description": item.get("description", ""),
                "site_id": item.get("site_id", ""),
                "vendor_id": item.get("vendor_id", "")
            }
            attrs = item.get("attributes", {}) or {}
            for col in dynamic_cols:
                key = col[5:]
                row[col] = attrs.get(key, "")
            rows.append(row)

        if file_format.lower() in ["xlsx", "excel"]:
            if not OPENPYXL_AVAILABLE:
                raise ValueError("Excel export requires 'openpyxl'.")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Items"
            ws.append(all_headers)
            for r in rows:
                ws.append([r.get(h, "") for h in all_headers])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "items_export.xlsx"
        else:
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=all_headers)
            writer.writeheader()
            for r in rows:
                writer.writerow(r)
            return output.getvalue().encode("utf-8-sig"), "text/csv", "items_export.csv"

