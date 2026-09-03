import io
import json
import openpyxl
import pytest

def test_template_download_csv(client):
    res = client.get("/api/v1/parts/import-template?format=csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["Content-Type"]
    assert "part_code" in res.data.decode("utf-8")
    assert "attr_resistance" in res.data.decode("utf-8") or "part_name" in res.data.decode("utf-8")

def test_template_download_xlsx(client):
    res = client.get("/api/v1/parts/import-template?format=xlsx")
    assert res.status_code == 200
    assert "openxmlformats" in res.headers["Content-Type"]
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Parts Catalog Template" in wb.sheetnames
    ws = wb["Parts Catalog Template"]
    assert ws.cell(row=1, column=1).value == "part_code"

def test_bulk_import_csv_with_auto_vendor_creation(client, admin_token):
    csv_data = """part_code,part_name,part_type,tracking_type,mpn,mfr,vendor,package,unit_of_measure,attr_resistance,attr_tolerance,attr_power_rating
TEST-RES-001,10K Resistor 0603,RESISTOR,QUANTITY,RC-001,Yageo,BrandNewVendor Inc,0603,PCS,10k,1%,0.1
TEST-RES-002,22K Resistor 0603,RESISTOR,QUANTITY,RC-002,Yageo,BrandNewVendor Inc,0603,PCS,22k,1%,0.1
"""
    data = {
        "file": (io.BytesIO(csv_data.encode("utf-8")), "parts.csv"),
        "duplicate_strategy": "skip"
    }
    res = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data=data,
        content_type="multipart/form-data"
    )
    assert res.status_code == 200
    res_data = res.json["data"]
    assert res_data["imported"] == 2
    assert "BrandNewVendor Inc" in res_data["auto_created_vendors"]
    assert len(res_data["errors"]) == 0

    # Verify part is queryable
    p_res = client.get("/api/v1/parts?search=TEST-RES-001")
    assert p_res.status_code == 200
    matched = [p for p in p_res.json["data"] if p["part_code"] == "TEST-RES-001"]
    assert len(matched) == 1
    assert matched[0]["vendor_name"] == "BrandNewVendor Inc"
    assert matched[0]["attributes"].get("resistance") == "10k"

def test_bulk_import_excel_format(client, admin_token):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["part_code", "part_name", "part_type", "tracking_type", "vendor", "attr_capacitance", "attr_voltage_rating"])
    ws.append(["TEST-CAP-XLSX", "100uF 25V Cap", "CAPACITOR", "QUANTITY", "ExcelVendor Corp", "100uF", "25V"])
    
    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    excel_buf.seek(0)

    data = {
        "file": (excel_buf, "parts.xlsx"),
        "duplicate_strategy": "skip"
    }
    res = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data=data,
        content_type="multipart/form-data"
    )
    assert res.status_code == 200
    res_data = res.json["data"]
    assert res_data["imported"] == 1
    assert "ExcelVendor Corp" in res_data["auto_created_vendors"]

def test_bulk_import_duplicate_strategy_skip_and_update(client, admin_token):
    # 1. First import
    csv_v1 = """part_code,part_name,part_type,tracking_type,description,attr_resistance,attr_tolerance,attr_power_rating
TEST-DUP-01,Original Name,RESISTOR,QUANTITY,Original Desc,100k,1%,0.125
"""
    res1 = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data={"file": (io.BytesIO(csv_v1.encode("utf-8")), "parts.csv"), "duplicate_strategy": "skip"},
        content_type="multipart/form-data"
    )
    assert res1.status_code == 200
    assert res1.json["data"]["imported"] == 1

    # 2. Re-import with 'skip'
    csv_v2 = """part_code,part_name,part_type,tracking_type,description,attr_resistance,attr_tolerance,attr_power_rating
TEST-DUP-01,Updated Name,RESISTOR,QUANTITY,Updated Desc,200k,1%,0.125
"""
    res_skip = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data={"file": (io.BytesIO(csv_v2.encode("utf-8")), "parts.csv"), "duplicate_strategy": "skip"},
        content_type="multipart/form-data"
    )
    assert res_skip.status_code == 200
    assert res_skip.json["data"]["imported"] == 0
    assert res_skip.json["data"]["skipped"] == 1

    # 3. Re-import with 'update'
    res_update = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data={"file": (io.BytesIO(csv_v2.encode("utf-8")), "parts.csv"), "duplicate_strategy": "update"},
        content_type="multipart/form-data"
    )
    assert res_update.status_code == 200
    assert res_update.json["data"]["updated"] == 1

    # Verify updated content
    p_res = client.get("/api/v1/parts?search=TEST-DUP-01")
    matched = [p for p in p_res.json["data"] if p["part_code"] == "TEST-DUP-01"]
    assert matched[0]["part_name"] == "Updated Name"
    assert matched[0]["description"] == "Updated Desc"

def test_bulk_import_json_format(client, admin_token):
    json_data = [
        {
            "part_code": "TEST-JSON-001",
            "part_name": "JSON Part",
            "part_type": "RESISTOR",
            "tracking_type": "QUANTITY",
            "vendor": "JsonVendor Ltd",
            "attr_resistance": "50k",
            "attr_tolerance": "1%",
            "attr_power_rating": "0.1"
        }
    ]
    res = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}", "Content-Type": "application/json"},
        json=json_data
    )
    assert res.status_code == 200
    assert res.json["data"]["imported"] == 1
    assert "JsonVendor Ltd" in res.json["data"]["auto_created_vendors"]

def test_bulk_import_validation_errors(client, admin_token):
    # CSV missing part name and invalid part type
    bad_csv = """part_code,part_name,part_type
,Missing Code,RESISTOR
BAD-CODE-01,,RESISTOR
BAD-CODE-02,Bad Type Part,NON_EXISTENT_TYPE
"""
    res = client.post(
        "/api/v1/parts/bulk-import",
        headers={"Authorization": f"Bearer {admin_token}"},
        data={"file": (io.BytesIO(bad_csv.encode("utf-8")), "bad.csv"), "duplicate_strategy": "skip"},
        content_type="multipart/form-data"
    )
    assert res.status_code == 200
    assert res.json["data"]["imported"] == 0
    assert len(res.json["data"]["errors"]) == 3
